import os
import tkinter as tk
from tkinter import Scrollbar, Entry, Button, messagebox, ttk
import threading
import subprocess
from watchdog.observers import Observer
from watchdog.events import FileSystemEventHandler
from docx import Document
from PyPDF2 import PdfReader
import openpyxl
import estilo
from folder_selector import seleccionar_carpeta

class FolderMonitorHandler(FileSystemEventHandler):
    def __init__(self, update_callback):
        self.update_callback = update_callback

    def on_modified(self, event):
        self.update_callback()

    def on_created(self, event):
        self.update_callback()

    def on_deleted(self, event):
        self.update_callback()

class FolderMonitorApp:
    def __init__(self, root, carpeta):
        self.root = root
        self.carpeta = carpeta
        self.root.title("Monitor de Carpeta con Buscador")
        self.root.configure(bg=estilo.COLOR_LISTADO)
        self.root.geometry(f"{estilo.ANCHO_VENTANA}x{estilo.ALTURA_VENTANA}")

        self.resultados_con_linea = {}
        self.exclusiones = ['node_modules', '.git', '.vs', '__pycache__']

        # Cabecera
        cabecera = tk.Frame(root, bg=estilo.COLOR_CABECERA)
        cabecera.pack(fill=tk.X)

        self.entry = Entry(cabecera, width=50, font=estilo.FUENTE_TEXTO)
        self.entry.pack(side=tk.LEFT, padx=10, pady=10)

        self.buscar_btn = Button(
            cabecera,
            text="Buscar por contenido",
            command=self.buscar,
            bg=estilo.COLOR_BOTON,
            fg=estilo.COLOR_TEXTO_BLANCO,
            font=estilo.FUENTE_TEXTO
        )
        self.buscar_btn.pack(side=tk.LEFT, padx=5)

        # Treeview
        self.tree = ttk.Treeview(root)
        self.tree.heading('#0', text='Archivos', anchor='w')
        self.tree.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)
        self.tree.bind("<Double-1>", self.abrir_archivo)
        self.tree.bind("<ButtonRelease-1>", self.previsualizar_linea)

        # Previsualización
        self.previsual_text = tk.Text(root, height=6, wrap='word', font=estilo.FUENTE_TEXTO)
        self.previsual_text.pack(fill=tk.X, padx=10, pady=(0, 10))
        self.previsual_text.configure(state='disabled')

        self.actualizar_lista()

        handler = FolderMonitorHandler(self.actualizar_lista)
        observer = Observer()
        observer.schedule(handler, self.carpeta, recursive=True)
        observer_thread = threading.Thread(target=observer.start)
        observer_thread.daemon = True
        observer_thread.start()

    def actualizar_lista(self):
        carpetas = []
        archivos = []

        for dirpath, dirnames, filenames in os.walk(self.carpeta):
            # Filtrar carpetas excluidas
            dirnames[:] = [d for d in dirnames if d not in self.exclusiones]

            rel_dir = os.path.relpath(dirpath, self.carpeta)
            if rel_dir == ".":
                rel_dir = ""

            for dirname in dirnames:
                ruta = os.path.normpath(os.path.join(rel_dir, dirname))
                carpetas.append(ruta)

            for filename in filenames:
                ruta = os.path.normpath(os.path.join(rel_dir, filename))
                if any(parte in self.exclusiones for parte in ruta.split(os.sep)):
                    continue
                archivos.append(ruta)

        self.todos_los_archivos = sorted(carpetas) + sorted(archivos)
        self.mostrar_lista(self.todos_los_archivos)

    def mostrar_lista(self, archivos):
        self.tree.delete(*self.tree.get_children())
        nodos = {}
        for archivo_relativo in archivos:
            partes = archivo_relativo.split(os.sep)
            padre = ''
            for i, parte in enumerate(partes):
                ruta_relativa = os.sep.join(partes[:i + 1])
                if ruta_relativa not in nodos:
                    nodo = self.tree.insert(padre, 'end', text=parte, open=True, values=[ruta_relativa])
                    nodos[ruta_relativa] = nodo
                padre = nodos[ruta_relativa]

    def buscar(self):
        palabra = self.entry.get().lower()
        self.resultados_con_linea.clear()

        if not palabra:
            self.mostrar_lista(self.todos_los_archivos)
            return

        encontrados = []
        for archivo_rel in self.todos_los_archivos:
            ruta = os.path.join(self.carpeta, archivo_rel)
            if not os.path.isfile(ruta):
                continue

            texto = ""
            nombre = archivo_rel.lower()

            try:
                if nombre.endswith(('.txt', '.md', '.csv', '.log', '.json', '.html', '.js', '.tsx', '.py', '.dtsx')):
                    with open(ruta, 'r', encoding='utf-8', errors='ignore') as f:
                        texto = f.read().lower()
                elif nombre.endswith('.docx'):
                    doc = Document(ruta)
                    texto = '\n'.join([p.text for p in doc.paragraphs]).lower()
                elif nombre.endswith('.pdf'):
                    reader = PdfReader(ruta)
                    texto = '\n'.join([page.extract_text() or "" for page in reader.pages]).lower()
                elif nombre.endswith('.xlsx'):
                    wb = openpyxl.load_workbook(ruta, data_only=True)
                    for sheet in wb.worksheets:
                        for row in sheet.iter_rows(values_only=True):
                            texto += ' '.join([str(cell) for cell in row if cell is not None]) + '\n'
                    texto = texto.lower()

                if palabra in texto:
                    for linea in texto.splitlines():
                        if palabra in linea:
                            self.resultados_con_linea[archivo_rel] = linea.strip()
                            break
                    encontrados.append(archivo_rel)

            except Exception as e:
                print(f"Error leyendo {archivo_rel}: {e}")

        self.mostrar_lista(encontrados)

    def previsualizar_linea(self, event):
        seleccion = self.tree.selection()
        if not seleccion:
            return
        ruta_rel = self.tree.item(seleccion[0])['values'][0]
        self.previsual_text.configure(state='normal')
        self.previsual_text.delete("1.0", tk.END)
        if ruta_rel in self.resultados_con_linea:
            self.previsual_text.insert(tk.END, self.resultados_con_linea[ruta_rel])
        self.previsual_text.configure(state='disabled')

    def abrir_archivo(self, event):
        seleccion = self.tree.selection()
        if seleccion:
            ruta_rel = self.tree.item(seleccion[0])['values'][0]
            ruta_completa = os.path.abspath(os.path.join(self.carpeta, ruta_rel))
            if os.path.isfile(ruta_completa):
                try:
                    if os.name == 'nt':
                        os.startfile(ruta_completa)
                    elif os.name == 'posix':
                        subprocess.call(('xdg-open', ruta_completa))
                except Exception as e:
                    messagebox.showerror("Error al abrir archivo", str(e))

if __name__ == "__main__":
    carpeta_seleccionada = seleccionar_carpeta()
    if carpeta_seleccionada:
        root = tk.Tk()
        app = FolderMonitorApp(root, carpeta_seleccionada)
        root.mainloop()
    else:
        print("No se seleccionó ninguna carpeta.")
